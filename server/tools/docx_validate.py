#!/usr/bin/env python3
"""
tool_docx_validate - Word 文档质量验证
"""
import json
import os
from pathlib import Path
from docx import Document

# ── Tool 定义 ────────────────────────────────────────

TOOL_NAME = "tool_docx_validate"
TOOL_DESCRIPTION = """验证 Word 文档质量和完整性，检查封面、内容、样式等。"""
TOOL_SCHEMA = {
    "type": "object",
    "properties": {
        "docx_path": {
            "type": "string",
            "description": "Word 文档路径"
        },
        "check_list": {
            "type": "array",
            "description": "检查项列表（可选，默认全部检查）",
            "items": {
                "type": "string",
                "enum": ["has_cover", "has_content", "min_size", "has_header", "has_footer", "has_tables", "headings"]
            }
        },
        "min_size_kb": {
            "type": "number",
            "description": "最小文件大小（KB），默认 10"
        }
    },
    "required": ["docx_path"]
}

# ── 检查函数 ────────────────────────────────────────

def validate_document(docx_path: str, check_list: list = None, min_size_kb: float = 10) -> dict:
    """验证文档"""
    try:
        if not os.path.exists(docx_path):
            return {"status": "FAILED", "error": f"文件不存在: {docx_path}"}

        doc = Document(docx_path)
        file_size_kb = os.path.getsize(docx_path) / 1024

        # 默认检查项
        all_checks = ["has_cover", "has_content", "min_size", "has_header", "has_footer", "has_tables", "headings"]
        checks_to_run = check_list if check_list else all_checks

        results = {}
        score = 0

        # has_cover: 段落数 > 5
        if "has_cover" in checks_to_run:
            results["has_cover"] = len(doc.paragraphs) > 5

        # has_content: 有非空段落
        if "has_content" in checks_to_run:
            results["has_content"] = any(p.text.strip() for p in doc.paragraphs)

        # min_size: 文件大小 >= 阈值
        if "min_size" in checks_to_run:
            results["min_size"] = file_size_kb >= min_size_kb

        # has_header: 有页眉
        if "has_header" in checks_to_run:
            if doc.sections:
                header_text = "".join([p.text for p in doc.sections[0].header.paragraphs])
                results["has_header"] = len(header_text.strip()) > 0 or len(doc.sections[0].header.tables) > 0
            else:
                results["has_header"] = False

        # has_footer: 有页脚
        if "has_footer" in checks_to_run:
            if doc.sections:
                footer_text = "".join([p.text for p in doc.sections[0].footer.paragraphs])
                results["has_footer"] = len(footer_text.strip()) > 0 or len(doc.sections[0].footer.tables) > 0
            else:
                results["has_footer"] = False

        # has_tables: 有表格
        if "has_tables" in checks_to_run:
            results["has_tables"] = len(doc.tables) > 0

        # headings: 标题统计
        if "headings" in checks_to_run:
            heading_counts = {}
            for para in doc.paragraphs:
                if para.style and para.style.name.startswith("Heading"):
                    level = para.style.name
                    heading_counts[level] = heading_counts.get(level, 0) + 1
            results["headings"] = {
                "total": sum(heading_counts.values()),
                "by_level": heading_counts
            }

        # 计算评分（布尔项平均值 * 100）
        bool_checks = [v for k, v in results.items() if isinstance(v, bool)]
        score = (sum(bool_checks) / len(bool_checks) * 100) if bool_checks else 100

        return {
            "status": "SUCCESS",
            "valid": all(bool_checks),
            "score": round(score, 1),
            "checks": results,
            "statistics": {
                "paragraphs": len(doc.paragraphs),
                "tables": len(doc.tables),
                "sections": len(doc.sections),
                "file_size_kb": round(file_size_kb, 1)
            }
        }

    except Exception as e:
        return {"status": "FAILED", "error": str(e), "valid": False, "score": 0}


# ── 主处理函数 ────────────────────────────────────────

async def handle(arguments: dict) -> str:
    """处理 tool_docx_validate 请求"""
    docx_path = arguments.get("docx_path", "")
    check_list = arguments.get("check_list")
    min_size_kb = arguments.get("min_size_kb", 10)

    if not docx_path:
        return json.dumps({"status": "FAILED", "error": "缺少 docx_path 参数"}, ensure_ascii=False)

    result = validate_document(docx_path, check_list, min_size_kb)
    return json.dumps(result, ensure_ascii=False, indent=2)


if __name__ == "__main__":
    import asyncio

    async def test():
        # 测试本地文件
        test_path = "C:/Users/Hnxz_/.openclaw/workspace/skills/Jason-Notion2Word/运营平台效率模块_V1.1.docx"
        result = await handle({"docx_path": test_path})
        print(result)

    asyncio.run(test())
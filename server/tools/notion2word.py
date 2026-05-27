#!/usr/bin/env python3
"""
tool_notion2word - Markdown 转 Word 文档
遵循 MD→Word 标准规则，支持自定义样式，支持图片嵌入
"""
import json
import os
import re
import tempfile
import requests
import io
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Tool 定义 ────────────────────────────────────────

TOOL_NAME = "tool_notion2word"
TOOL_DESCRIPTION = """将 Markdown 内容转换为 Word 文档。支持标题、段落、列表、表格、代码块等元素。
输出文件保存在 /app/files/ 目录下。"""
TOOL_SCHEMA = {
    "type": "object",
    "properties": {
        "markdown": {
            "type": "string",
            "description": "Markdown 格式的内容"
        },
        "output": {
            "type": "string",
            "description": "输出 Word 文件路径（相对于 /app/files/）"
        },
        "title": {
            "type": "string",
            "description": "文档标题（可选，默认从内容提取）"
        },
        "rules": {
            "type": "object",
            "description": "自定义样式规则（可选）",
            "properties": {
                "font_name": {"type": "string", "description": "字体名称，默认宋体"},
                "heading_sizes": {"type": "object", "description": "标题字号配置"},
                "body_size": {"type": "number", "description": "正文字号，默认10.5"}
            }
        }
    },
    "required": ["markdown", "output"]
}

# ── 默认样式规则 ────────────────────────────────────────

DEFAULT_RULES = {
    "font_name": "宋体",
    "heading_sizes": {
        1: 16,  # 三号
        2: 14,  # 四号
        3: 12,  # 小四
        4: 12   # 小四
    },
    "body_size": 10.5,  # 五号
    "code_font": "宋体",  # 代码块也用宋体
    "code_size": 10,
    "line_spacing": 1.5,
    "first_line_indent": 0.74,  # 2字符（cm）
    "font_color": "000000"  # 黑色
}

# ── 辅助函数 ────────────────────────────────────────

def set_chinese_font(run, font_name='宋体', size_pt=None):
    """设置中文字体（黑色）"""
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    if size_pt:
        run.font.size = Pt(size_pt)
    # 强制设置字体颜色为黑色
    run.font.color.rgb = RGBColor(0, 0, 0)

def create_styles(doc, rules):
    """创建文档样式"""
    styles = doc.styles
    font_name = rules.get('font_name', '宋体')
    heading_sizes = rules.get('heading_sizes', DEFAULT_RULES['heading_sizes'])
    body_size = rules.get('body_size', DEFAULT_RULES['body_size'])

    # 标题样式
    for level in range(1, 5):
        style_name = f'Heading {level}'
        try:
            style = styles[style_name]
        except KeyError:
            style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)

        style.font.name = font_name
        style.font.size = Pt(heading_sizes.get(level, 12))
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)  # 黑色
        style._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    # 正文样式
    try:
        normal_style = styles['Normal']
    except KeyError:
        normal_style = styles.add_style('Normal', WD_STYLE_TYPE.PARAGRAPH)

    normal_style.font.name = font_name
    normal_style.font.size = Pt(body_size)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)  # 黑色
    normal_style._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    normal_style.paragraph_format.first_line_indent = Cm(rules.get('first_line_indent', 0.74))
    normal_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    return doc

def parse_table(lines):
    """解析 Markdown 表格"""
    rows = []
    for line in lines:
        if line.startswith('|') and line.endswith('|'):
            cells = [cell.strip() for cell in line[1:-1].split('|')]
            # 跳过分隔行
            if all(set(cell) <= {'-', ':', ' '} for cell in cells):
                continue
            rows.append(cells)
    return rows

def add_table(doc, rows, rules):
    """添加表格到文档"""
    if not rows:
        return

    font_name = rules.get('font_name', '宋体')
    body_size = rules.get('body_size', 10.5)

    num_cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=num_cols)
    table.style = 'Table Grid'

    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            if j < num_cols:
                cell = table.rows[i].cells[j]
                cell.text = ""
                para = cell.paragraphs[0]
                run = para.add_run(cell_text)
                set_chinese_font(run, font_name, body_size)

def download_image(url, timeout=30):
    """下载图片并返回字节流，支持URL和本地文件路径"""
    try:
        # 检查是否是本地文件路径
        if url.startswith('/') or url.startswith('C:') or url.startswith('D:'):
            if os.path.exists(url):
                with open(url, 'rb') as f:
                    return io.BytesIO(f.read())
            else:
                print(f"[Image] Local file not found: {url}")
                return None

        # 远程URL下载
        headers = {
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'
        }
        response = requests.get(url, headers=headers, timeout=timeout, stream=True)
        if response.status_code == 200:
            return io.BytesIO(response.content)
    except Exception as e:
        print(f"[Image] Download failed: {url[:50]}... - {e}")
    return None

def add_image(doc, url, alt_text="", max_width_inches=6):
    """添加图片到文档"""
    image_stream = download_image(url)
    if image_stream:
        try:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            # 添加图片，设置最大宽度
            picture = run.add_picture(image_stream)
            # 限制图片宽度
            if picture.width > Inches(max_width_inches):
                ratio = Inches(max_width_inches) / picture.width
                picture.width = Inches(max_width_inches)
                picture.height = int(picture.height * ratio)
            # 添加图片说明
            if alt_text:
                caption_p = doc.add_paragraph()
                caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption_run = caption_p.add_run(alt_text)
                caption_run.font.size = Pt(9)
                caption_run.font.color.rgb = RGBColor(128, 128, 128)
            return True
        except Exception as e:
            print(f"[Image] Insert failed: {e}")
    # 图片下载失败，插入占位符
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"[图片: {alt_text or url[:50]}]")
    run.font.color.rgb = RGBColor(128, 128, 128)
    run.font.size = Pt(10)
    return False

def parse_markdown(markdown, doc, rules):
    """解析 Markdown 并添加到文档"""
    lines = markdown.split('\n')
    font_name = rules.get('font_name', '宋体')
    body_size = rules.get('body_size', 10.5)
    code_font = rules.get('code_font', '宋体')
    code_size = rules.get('code_size', 10)
    first_line_indent = rules.get('first_line_indent', 0.74)

    i = 0
    in_code_block = False
    code_lines = []
    code_lang = ""
    in_table = False
    table_lines = []

    while i < len(lines):
        line = lines[i]

        # 代码块处理
        if line.startswith('```'):
            if in_code_block:
                # 结束代码块
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Cm(0)
                code_text = '\n'.join(code_lines)
                run = p.add_run(code_text)
                set_chinese_font(run, font_name, code_size)  # 使用宋体
                code_lines = []
                in_code_block = False
            else:
                # 开始代码块
                in_code_block = True
                code_lang = line[3:].strip()
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # 表格处理
        if line.startswith('|'):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            # 表格结束
            rows = parse_table(table_lines)
            add_table(doc, rows, rules)
            doc.add_paragraph()  # 表格后空行
            in_table = False
            table_lines = []

        # 标题
        heading_match = re.match(r'^(#{1,4})\s+(.+)$', line)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            p = doc.add_heading(level=level)
            run = p.add_run(text)
            set_chinese_font(run, font_name, rules.get('heading_sizes', {}).get(level, 12))
            run.bold = True
            i += 1
            continue

        # 无序列表
        bullet_match = re.match(r'^[-*]\s+(.+)$', line)
        if bullet_match:
            text = bullet_match.group(1)
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(text)
            set_chinese_font(run, font_name, body_size)
            p.paragraph_format.first_line_indent = Cm(0)
            i += 1
            continue

        # 有序列表
        numbered_match = re.match(r'^\d+\.\s+(.+)$', line)
        if numbered_match:
            text = numbered_match.group(1)
            p = doc.add_paragraph(style='List Number')
            run = p.add_run(text)
            set_chinese_font(run, font_name, body_size)
            p.paragraph_format.first_line_indent = Cm(0)
            i += 1
            continue

        # 引用
        quote_match = re.match(r'^>\s+(.+)$', line)
        if quote_match:
            text = quote_match.group(1)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            run = p.add_run(text)
            set_chinese_font(run, font_name, body_size)
            i += 1
            continue

        # 图片处理 ![alt](url)
        image_match = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)$', line.strip())
        if image_match:
            alt_text = image_match.group(1)
            image_url = image_match.group(2)
            add_image(doc, image_url, alt_text)
            i += 1
            continue

        # 行内图片处理（段落中包含图片）
        inline_image_pattern = r'!\[([^\]]*)\]\(([^)]+)\)'
        if re.search(inline_image_pattern, line):
            # 提取所有图片
            parts = re.split(inline_image_pattern, line)
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0)
            idx = 0
            while idx < len(parts):
                if idx % 3 == 0:
                    # 文本部分
                    if parts[idx].strip():
                        run = p.add_run(parts[idx])
                        set_chinese_font(run, font_name, body_size)
                elif idx % 3 == 1:
                    # alt text
                    pass
                elif idx % 3 == 2:
                    # URL - 下载并插入图片
                    img_url = parts[idx]
                    img_stream = download_image(img_url)
                    if img_stream:
                        try:
                            run = p.add_run()
                            run.add_picture(img_stream, width=Inches(2))
                        except:
                            run = p.add_run(f"[图]")
                            run.font.color.rgb = RGBColor(128, 128, 128)
                idx += 1
            i += 1
            continue

        # 分隔线
        if re.match(r'^-{3,}$', line.strip()):
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0)
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), 'auto')
            pBdr.append(bottom)
            p._p.get_or_add_pPr().append(pBdr)
            i += 1
            continue

        # 普通段落
        if line.strip():
            p = doc.add_paragraph()
            run = p.add_run(line)
            set_chinese_font(run, font_name, body_size)
            p.paragraph_format.first_line_indent = Cm(first_line_indent)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

        i += 1

    # 处理未结束的表格
    if in_table and table_lines:
        rows = parse_table(table_lines)
        add_table(doc, rows, rules)

# ── 主处理函数 ────────────────────────────────────────

async def handle(arguments: dict) -> str:
    """处理 tool_notion2word 请求"""
    markdown = arguments.get('markdown', '')
    output = arguments.get('output', '')
    title = arguments.get('title', '')
    custom_rules = arguments.get('rules', {})

    if not markdown:
        return json.dumps({"status": "FAILED", "error": "markdown 参数为空"}, ensure_ascii=False)

    if not output:
        return json.dumps({"status": "FAILED", "error": "output 参数为空"}, ensure_ascii=False)

    # 合并规则
    rules = {**DEFAULT_RULES, **custom_rules}

    # 输出路径处理 - 支持绝对路径和相对路径
    if output.startswith('/') or (len(output) > 2 and output[1] == ':'):
        # 绝对路径
        output_path = Path(output)
    else:
        # 相对路径（兼容Docker环境和本地环境）
        if os.path.exists('/app/files'):
            output_path = Path('/app/files') / output
        else:
            output_path = Path(output)

    try:
        # 创建文档
        doc = Document()
        doc = create_styles(doc, rules)

        # 添加标题
        if title:
            title_p = doc.add_heading(level=0)
            title_run = title_p.add_run(title)
            set_chinese_font(title_run, rules.get('font_name', '宋体'), 18)
            title_run.bold = True
            title_run.font.color.rgb = RGBColor(0, 0, 0)  # 黑色
            title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 解析 Markdown
        parse_markdown(markdown, doc, rules)

        # 确保输出目录存在
        output_path.parent.mkdir(parents=True, exist_ok=True)

        # 保存文档
        doc.save(str(output_path))

        return json.dumps({
            "status": "SUCCESS",
            "message": f"Word 文档已生成",
            "output": str(output_path),
            "size_bytes": output_path.stat().st_size
        }, ensure_ascii=False)

    except Exception as e:
        return json.dumps({"status": "FAILED", "error": str(e)}, ensure_ascii=False)